"""Downloadable files (the `bulk_file` connector).

The fetch layer finds the file one of three ways: a fixed `url`, a link `discover`ed on a page
(for files whose URL carries a hash that changes, like the Pink Sheet), or a `dataverse`
dataset (PWT, Maddison). Every download must be the expected file type by its leading bytes:
an HTML answer, which is how a bot challenge arrives, fails the run and is never worked around.
Each dialect then parses its own file layout into the common time-series model.
"""

import csv
import gzip
import hashlib
import io
import json
import re
from collections import Counter
from collections.abc import Callable, Iterator, Mapping
from dataclasses import dataclass
from datetime import date
from typing import Any, Literal
from urllib.parse import unquote, urlparse

import httpx
import openpyxl
from pydantic import BaseModel, ConfigDict, Field, model_validator

from dome_open_data.connectors import Collected, RawFile, SeriesMeta
from dome_open_data.http import RateLimiter, SourceError, get, require_file, require_json
from dome_open_data.model import Observation
from dome_open_data.registry import Source

BULK_DIALECTS = {"pinksheet", "wpp", "pwt", "maddison"}

_PINK_PERIOD = re.compile(r"^(\d{4})M(0[1-9]|1[0-2])$")
_TRAILING_PARENTHETICAL = re.compile(r"^(.*?)\s*\(([^()]*)\)\s*$")
# openpyxl's read-only mode trusts the sheet's declared size, which can claim a million rows.
_EMPTY_ROWS_BEFORE_STOP = 100


class SeriesSpec(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    title: str = Field(min_length=1)
    unit: str


class Discover(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    page: str = Field(pattern=r"^https://")
    pattern: str = Field(min_length=1)


class Dataverse(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    doi: str = Field(pattern=r"^doi:")
    filename: str = Field(min_length=1)


class BulkOptions(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    url: str | None = Field(default=None, pattern=r"^https://")
    discover: Discover | None = None
    dataverse: Dataverse | None = None
    file_type: Literal["xlsx", "gzip"]
    sheet: str | None = None
    series: dict[str, SeriesSpec] = Field(default_factory=dict)
    exclude: list[str] = Field(default_factory=list)
    loc_types: list[str] = Field(default_factory=list)
    projection_from: int | None = None
    geo: str | None = None

    @model_validator(mode="after")
    def _one_location(self) -> "BulkOptions":
        if sum(x is not None for x in (self.url, self.discover, self.dataverse)) != 1:
            raise ValueError("exactly one of url, discover, dataverse")
        return self


@dataclass(frozen=True, slots=True)
class Resolved:
    url: str
    name: str
    token: str
    source_url: str


@dataclass
class Parsed:
    observations: list[Observation]
    series: list[SeriesMeta]
    rows_raw: int
    dropped: dict[str, int]


def _is_number(value: Any) -> bool:
    return isinstance(value, int | float) and not isinstance(value, bool)


def _rows(workbook: Any, sheet: str) -> Iterator[tuple[Any, ...]]:
    """Rows of a sheet, stopping after a long run of empty rows."""
    if sheet not in workbook.sheetnames:
        raise SourceError(f"sheet {sheet!r} not found; sheets are {workbook.sheetnames}")
    empty = 0
    for row in workbook[sheet].iter_rows(values_only=True):
        if all(cell is None or cell == "" for cell in row):
            empty += 1
            if empty >= _EMPTY_ROWS_BEFORE_STOP:
                return
            continue
        empty = 0
        yield row


def _workbook(data: bytes) -> Any:
    return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def slug(text: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", text.upper()).strip("_")


def split_definition(text: str) -> tuple[str, str]:
    """'Population (in millions)' -> ('Population', 'millions')."""
    match = _TRAILING_PARENTHETICAL.match(text.strip())
    if not match:
        return text.strip(), ""
    unit = match[2].strip()
    return match[1].strip(), unit[3:] if unit.startswith("in ") else unit


class BulkFile:
    def __init__(self, source: Source) -> None:
        self.source = source
        self.options = BulkOptions.model_validate(source.options)
        self.limiter = RateLimiter(source.rate_limit.per_second)
        self.source_url = ""
        parsers: dict[str, Callable[[bytes], Parsed]] = {
            "pinksheet": self._pinksheet,
            "wpp": self._wpp,
            "pwt": self._pwt,
            "maddison": self._maddison,
        }
        self.parse = parsers[str(source.dialect)]

    # Fetch

    def _selection(self) -> str:
        options = self.options.model_dump(mode="json", exclude={"url", "discover", "dataverse"})
        return hashlib.sha256(json.dumps(options, sort_keys=True).encode()).hexdigest()[:12]

    async def _header_signal(self, client: httpx.AsyncClient, url: str) -> str:
        response = await get(client, url, self.limiter, method="HEAD")
        name = "etag" if self.source.change_signal == "etag" else "last-modified"
        value: str = response.headers.get(name, "")
        if not value:
            raise SourceError(f"HEAD {url}: no {name} header to detect changes")
        return value

    async def resolve(self, client: httpx.AsyncClient) -> Resolved:
        options = self.options
        if options.url:
            url = options.url
            signal = await self._header_signal(client, url)
            name = unquote(urlparse(url).path.rsplit("/", 1)[-1])
            return Resolved(url, name, f"{url}|{signal}|{self._selection()}", url)
        if options.discover:
            page = await get(client, options.discover.page, self.limiter)
            links = sorted(set(re.findall(options.discover.pattern, page.text)))
            if len(links) != 1:
                raise SourceError(
                    f"{options.discover.page}: expected one link matching "
                    f"{options.discover.pattern!r}, found {links}"
                )
            url = links[0]
            signal = await self._header_signal(client, url)
            name = unquote(urlparse(url).path.rsplit("/", 1)[-1])
            token = f"{url}|{signal}|{self._selection()}"
            return Resolved(url, name, token, options.discover.page)
        assert options.dataverse is not None
        dataverse = options.dataverse
        api = f"{self.source.endpoint}api/datasets/:persistentId/"
        response = await get(client, api, self.limiter, params={"persistentId": dataverse.doi})
        version = require_json(response)["data"]["latestVersion"]
        files = [f["dataFile"] for f in version["files"]]
        match = [f for f in files if f["filename"] == dataverse.filename]
        if len(match) != 1:
            raise SourceError(f"{dataverse.doi}: no single file named {dataverse.filename}")
        file = match[0]
        checksum = (file.get("checksum") or {}).get("value") or file.get("md5") or ""
        token = (
            f"{dataverse.doi}|v{version['versionNumber']}.{version['versionMinorNumber']}"
            f"|{file['id']}|{checksum}|{self._selection()}"
        )
        url = f"{self.source.endpoint}api/access/datafile/{file['id']}"
        doi_url = f"https://doi.org/{dataverse.doi.removeprefix('doi:')}"
        return Resolved(url, dataverse.filename, token, doi_url)

    async def change_token(self, client: httpx.AsyncClient, source: Source) -> str:
        return (await self.resolve(client)).token

    async def collect(self, client: httpx.AsyncClient, source: Source, vintage: date) -> Collected:
        resolved = await self.resolve(client)
        self.source_url = resolved.source_url
        response = await get(client, resolved.url, self.limiter)
        data = require_file(response, self.options.file_type)
        parsed = self.parse(data)
        return Collected(
            raw_files=[RawFile(resolved.name, resolved.url, data)],
            observations=parsed.observations,
            series=parsed.series,
            rows_raw=parsed.rows_raw,
            dropped=parsed.dropped,
        )

    def count_clean_rows(self, files: Mapping[str, bytes]) -> int:
        data = [content for name, content in files.items() if not name.startswith("_")]
        if len(data) != 1:
            raise SourceError(f"{self.source.id}: expected one raw file, found {len(data)}")
        return len(self.parse(data[0]).observations)

    # Parse

    def _builder(self) -> "_Builder":
        return _Builder(self.source, self.source_url)

    def _pinksheet(self, data: bytes) -> Parsed:
        rows = _rows(_workbook(data), self.options.sheet or "Monthly Prices")
        names: tuple[Any, ...] | None = None
        units: tuple[Any, ...] | None = None
        for row in rows:
            if names is None:
                if len(row) > 1 and row[0] is None and isinstance(row[1], str) and row[1].strip():
                    names = row
            else:
                units = row
                break
        if names is None or units is None:
            raise SourceError("Pink Sheet: no header rows (names, then units) found")
        columns: list[tuple[int, str, str, str]] = []
        for i, name in enumerate(names):
            if i == 0 or not isinstance(name, str) or not name.strip():
                continue
            title = name.replace("*", "").strip()
            unit = str(units[i] or "").strip().removeprefix("(").removesuffix(")")
            columns.append((i, slug(title), title, unit))
        duplicates = [c for c, n in Counter(code for _, code, _, _ in columns).items() if n > 1]
        if duplicates:
            raise SourceError(f"Pink Sheet: duplicate series codes {duplicates}")
        build = self._builder()
        geo = self.options.geo or "WLD"
        for row in rows:
            match = _PINK_PERIOD.match(str(row[0] or "").strip())
            if not match:
                continue
            period = date(int(match[1]), int(match[2]), 1)
            for i, code, title, unit in columns:
                build.add(code, geo, period, "M", row[i] if i < len(row) else None, title, unit)
        return build.parsed()

    def _wpp(self, data: bytes) -> Parsed:
        series = self.options.series
        if not series or not self.options.loc_types:
            raise SourceError("WPP: options.series and options.loc_types are required")
        reader = csv.DictReader(io.StringIO(gzip.decompress(data).decode("utf-8-sig")))
        needed = {"LocID", "ISO3_code", "LocTypeName", "Time", *series}
        missing = needed - set(reader.fieldnames or [])
        if missing:
            raise SourceError(f"WPP: columns missing {sorted(missing)}")
        build = self._builder()
        loc_types = set(self.options.loc_types)
        for row in reader:
            if row["LocTypeName"] not in loc_types:
                continue
            geo = row["ISO3_code"] or f"UN{row['LocID']}"
            year = int(row["Time"])
            status = (
                "projection"
                if self.options.projection_from and year >= self.options.projection_from
                else None
            )
            for code, spec in series.items():
                cell = row[code]
                value = float(cell) if cell != "" else None
                build.add(code, geo, date(year, 1, 1), "A", value, spec.title, spec.unit, status)
        return build.parsed()

    def _pwt(self, data: bytes) -> Parsed:
        workbook = _workbook(data)
        legend = {
            row[0]: row[1]
            for row in _rows(workbook, "Legend")
            if isinstance(row[0], str) and isinstance(row[1], str)
        }
        rows = _rows(workbook, self.options.sheet or "Data")
        header = next(rows)
        index = {name: i for i, name in enumerate(header) if isinstance(name, str)}
        for required in ("countrycode", "year"):
            if required not in index:
                raise SourceError(f"PWT: no {required!r} column")
        identifiers = {"countrycode", "country", "currency_unit", "year"}
        specs: dict[str, SeriesSpec] = {}
        for name in index:
            if name in identifiers or name in self.options.exclude:
                continue
            if name in self.options.series:
                specs[name] = self.options.series[name]
            elif name in legend:
                title, unit = split_definition(legend[name])
                specs[name] = SeriesSpec(title=title, unit=unit)
            else:
                raise SourceError(f"PWT: variable {name!r} has no definition in the legend")
        build = self._builder()
        for row in rows:
            geo, year = row[index["countrycode"]], row[index["year"]]
            for name, spec in specs.items():
                period = date(int(year), 1, 1) if _is_number(year) else None
                build.add(name, geo, period, "A", row[index[name]], spec.title, spec.unit)
        return build.parsed()

    def _maddison(self, data: bytes) -> Parsed:
        series = self.options.series
        if not series:
            raise SourceError("Maddison: options.series is required")
        rows = _rows(_workbook(data), self.options.sheet or "Full data")
        header = next(rows)
        index = {name: i for i, name in enumerate(header) if isinstance(name, str)}
        missing = {"countrycode", "year", *series} - index.keys()
        if missing:
            raise SourceError(f"Maddison: columns missing {sorted(missing)}")
        build = self._builder()
        for row in rows:
            geo, year = row[index["countrycode"]], row[index["year"]]
            period = date(int(year), 1, 1) if _is_number(year) else None
            for code, spec in series.items():
                build.add(code, geo, period, "A", row[index[code]], spec.title, spec.unit)
        return build.parsed()


class _Builder:
    """Collects observations and series metadata, counting every candidate cell."""

    def __init__(self, source: Source, source_url: str) -> None:
        self.source = source
        self.source_url = source_url
        self.observations: list[Observation] = []
        self.series: dict[str, SeriesMeta] = {}
        self.rows_raw = 0
        self.dropped: Counter[str] = Counter()

    def add(
        self,
        code: str,
        geo: Any,
        period: date | None,
        freq: str,
        value: Any,
        title: str,
        unit: str,
        status: str | None = None,
    ) -> None:
        self.rows_raw += 1
        if not _is_number(value):
            self.dropped["null_value"] += 1
            return
        if not isinstance(geo, str) or not geo.strip():
            self.dropped["no_geo"] += 1
            return
        if period is None:
            self.dropped["no_period"] += 1
            return
        geo = geo.strip()
        series_id = f"{self.source.publisher}.{self.source.dataset}.{code}.{geo}.{freq}"
        self.observations.append(
            Observation(series_id, geo, period, freq, float(value), unit, status)
        )
        if series_id not in self.series:
            self.series[series_id] = SeriesMeta(
                series_id=series_id,
                code=code,
                geo=geo,
                freq=freq,
                title=title,
                definition=None,
                unit=unit,
                seasonal_adjustment=None,
                source_url=self.source_url,
            )

    def parsed(self) -> Parsed:
        return Parsed(
            self.observations, list(self.series.values()), self.rows_raw, dict(self.dropped)
        )
